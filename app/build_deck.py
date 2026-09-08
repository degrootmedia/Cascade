# -*- coding: utf-8 -*-
"""Generate Cascade_Pitch.pptx — the full internal-tool pitch deck (v2).

15 slides, 16:9. Part 1 = Operations (centralize, replace Boords, expense ledger).
Part 2 = Creative toolkit (pipeline, script/style, boards, motion tools, handoff).
Framing: Bring Your Own API Key — supports OpenArt and Higgsfield. Never mentions
the underlying chat vendor. Editable in Canva: import the .pptx at canva.com
(Create a design > Import file).
"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

INK     = RGBColor(0x1B, 0x17, 0x40)
INK2    = RGBColor(0x27, 0x21, 0x5C)
CREAM   = RGBColor(0xFD, 0xF8, 0xEF)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CORAL   = RGBColor(0xFF, 0x6B, 0x5E)
AMBER   = RGBColor(0xFF, 0xC1, 0x45)
TEAL    = RGBColor(0x2E, 0xC4, 0xB6)
VIOLET  = RGBColor(0x7C, 0x5C, 0xFF)
PINK    = RGBColor(0xFF, 0x5D, 0x8F)
LILAC   = RGBColor(0xC9, 0xC2, 0xFF)
GRAY    = RGBColor(0x5E, 0x5A, 0x78)
DARKTXT = RGBColor(0x22, 0x1E, 0x4E)

FONT = "Trebuchet MS"
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = INK if dark else CREAM
    return s


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12, dash=False, line_w=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1.5)
        if dash:
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.0):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing != 1.0:
            p.line_spacing = line_spacing
        p.space_after = Pt(0)
        for (t, size, color, bold, *rest) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = rest[0] if rest else FONT
    return tb


def P(t, size, color, bold=False, font=FONT):
    return (t, size, color, bold, font)


def chip(s, x, y, w, h, label, fill, color, size=13, bold=True, line=None):
    sp = box(s, x, y, w, h, fill=fill, line=line, radius=0.5)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(10); tf.margin_right = Pt(10)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = FONT
    return sp


def kicker(s, x, y, label, color):
    text(s, x, y, Inches(7), Inches(0.4), [[P(label.upper(), 13, color, True)]])


def placeholder(s, x, y, w, h, label, dark=True):
    fill = INK2 if dark else WHITE
    edge = LILAC if dark else GRAY
    txtc = LILAC if dark else GRAY
    box(s, x, y, w, h, fill=fill, line=edge, dash=True, line_w=Pt(1.75), radius=0.06)
    text(s, x, y + h / 2 - Inches(0.55), w, Inches(1.1),
         [[P("SCREENSHOT", 16, txtc, True)], [P(label, 12, txtc, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def dots(s, x, y, colors, r=Inches(0.16), gap=Inches(0.30)):
    for i, c in enumerate(colors):
        box(s, x + Emu(int(gap * i)), y, r, r, fill=c, shape=MSO_SHAPE.OVAL)


def footer(s, n, label, dark=True):
    c = LILAC if dark else GRAY
    text(s, Inches(0.55), Inches(7.02), Inches(8), Inches(0.35),
         [[P(label, 10, c, False)]])
    text(s, Inches(12.35), Inches(7.02), Inches(0.6), Inches(0.35),
         [[P("%02d" % n, 10, c, True)]], align=PP_ALIGN.RIGHT)


def blob_strip(s, dark=True):
    colors = [CORAL, AMBER, TEAL, VIOLET, PINK]
    dots(s, Inches(0.55), Inches(0.45), colors)


def bullet_list(s, x, y, w, items, dark=True, size=14, gap=0.62, bullet_color=None):
    c = WHITE if dark else DARKTXT
    bc = bullet_color or AMBER
    for i, it in enumerate(items):
        box(s, x, y + Emu(int(Inches(gap) * i)) + Inches(0.07), Inches(0.14),
            Inches(0.14), fill=bc, shape=MSO_SHAPE.OVAL)
        text(s, x + Inches(0.32), y + Emu(int(Inches(gap) * i)), w - Inches(0.32),
             Inches(0.55), [[P(it, size, c, False)]], line_spacing=1.05)


def card(s, x, y, w, h, top_color, title, body, dark=False, title_size=15,
         body_size=12.5, icon=None):
    bg = INK2 if dark else WHITE
    box(s, x, y, w, h, fill=bg, radius=0.10)
    box(s, x, y, w, Inches(0.14), fill=top_color, radius=0.5)
    head = []
    if icon:
        head.append(P(icon + "  ", title_size, top_color, True))
    head.append(P(title, title_size, WHITE if dark else DARKTXT, True))
    text(s, x + Inches(0.28), y + Inches(0.32), w - Inches(0.56), Inches(0.5), [head])
    text(s, x + Inches(0.28), y + Inches(0.88), w - Inches(0.56),
         h - Inches(1.05), [[P(body, body_size, LILAC if dark else GRAY, False)]],
         line_spacing=1.08)


R = "\u2192"

# ---------- SLIDE 1 : TITLE ----------
s = slide(dark=True)
box(s, Inches(10.4), Inches(-1.8), Inches(5.6), Inches(5.6), fill=VIOLET, shape=MSO_SHAPE.OVAL)
box(s, Inches(11.6), Inches(-0.9), Inches(4.4), Inches(4.4), fill=TEAL, shape=MSO_SHAPE.OVAL)
box(s, Inches(-1.9), Inches(5.2), Inches(4.6), Inches(4.6), fill=CORAL, shape=MSO_SHAPE.OVAL)
box(s, Inches(-1.1), Inches(6.2), Inches(3.0), Inches(3.0), fill=AMBER, shape=MSO_SHAPE.OVAL)
text(s, Inches(0.9), Inches(1.55), Inches(7), Inches(0.4),
     [[P("INTERNAL TOOL PROPOSAL  \u00b7  PRODUCTION OPERATIONS + CREATIVE PROCESS", 13, AMBER, True)]])
text(s, Inches(0.85), Inches(1.95), Inches(10.5), Inches(2.3),
     [[P("Cascade", 96, WHITE, True)]])
text(s, Inches(0.9), Inches(3.8), Inches(9.4), Inches(1.4),
     [[P("One studio desktop for every AI step of production \u2014 from first draft to editor handoff, on our own API keys.", 22, LILAC, False)]],
     line_spacing=1.2)
chip(s, Inches(0.9), Inches(5.35), Inches(2.95), Inches(0.52), "Bring Your Own API Key", TEAL, INK, 12.5)
chip(s, Inches(4.05), Inches(5.35), Inches(1.75), Inches(0.52), "OpenArt", VIOLET, WHITE, 12.5)
chip(s, Inches(5.95), Inches(5.35), Inches(2.05), Inches(0.52), "Higgsfield", PINK, WHITE, 12.5)
chip(s, Inches(8.15), Inches(5.35), Inches(2.45), Inches(0.52), "No new subscriptions", AMBER, INK, 12.5)
footer(s, 1, "Cascade \u00b7 Internal pitch", dark=True)

# ---------- SLIDE 2 : THE PROBLEM ----------
s = slide(dark=False)
blob_strip(s)
kicker(s, Inches(0.55), Inches(0.95), "The problem", CORAL)
text(s, Inches(0.55), Inches(1.3), Inches(11.5), Inches(1.5),
     [[P("AI is everywhere in our pipeline \u2014 ", 32, DARKTXT, True),
       P("and nowhere we can see it.", 32, CORAL, True)]], line_spacing=1.1)
cw = Inches(2.92); gapx = Inches(0.18); y0 = Inches(3.1); ch = Inches(3.15)
card(s, Inches(0.55), y0, cw, ch, CORAL, "Scattered everywhere",
     "Every artist brings their own accounts and browser tabs. Output lives in downloads folders, chat threads, and email.", dark=False, icon="S:")
card(s, Inches(0.55) + cw + gapx, y0, cw, ch, AMBER, "Zero cost visibility",
     "No one can answer what AI cost us this month. Spend is split across personal cards and mystery invoices.", dark=False, icon="$$")
card(s, Inches(0.55) + 2 * (cw + gapx), y0, cw, ch, VIOLET, "Subscription sprawl",
     "We pay monthly for point tools that each cover one small slice of the job.", dark=False, icon="X X")
card(s, Inches(0.55) + 3 * (cw + gapx), y0, cw, ch, TEAL, "Broken handoffs",
     "Scripts, boards, and animatic files are stitched together by hand and passed around by folder-dumping.", dark=False, icon=">>")
footer(s, 2, "Why now", dark=False)

# ---------- SLIDE 3 : MEET CASCADE ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "The proposal", TEAL)
text(s, Inches(0.55), Inches(0.9), Inches(11), Inches(0.9),
     [[P("Meet ", 34, WHITE, True), P("Cascade", 34, TEAL, True),
       P(" \u2014 our own AI production desk", 34, WHITE, True)]])
text(s, Inches(0.55), Inches(1.8), Inches(5.6), Inches(1.2),
     [[P("A single Windows desktop app the whole team runs. It centralizes every AI workflow \u2014 writing, boards, animatics, and delivery \u2014 and connects to the generation services we already use, on our own API keys.", 15, LILAC, False)]],
     line_spacing=1.15)
cy = Inches(3.1)
chip(s, Inches(0.55), cy, Inches(5.35), Inches(0.5), "Built-in AI agent \u2014 chat, skills & custom agents", INK2, LILAC, 12.5, line=LILAC)
chip(s, Inches(0.55), cy + Inches(0.62), Inches(5.35), Inches(0.5), "Image generation via OpenArt & Higgsfield", INK2, LILAC, 12.5, line=LILAC)
chip(s, Inches(0.55), cy + Inches(1.24), Inches(5.35), Inches(0.5), "Script " + R + " boards " + R + " animatic " + R + " editor handoff", INK2, LILAC, 12.5, line=LILAC)
chip(s, Inches(0.55), cy + Inches(1.86), Inches(5.35), Inches(0.5), "Bring Your Own API Key \u2014 no middleman markup", INK2, LILAC, 12.5, line=LILAC)
chip(s, Inches(0.55), cy + Inches(2.48), Inches(5.35), Inches(0.5), "Every generation logged to a studio expense ledger", INK2, LILAC, 12.5, line=LILAC)
placeholder(s, Inches(6.35), Inches(1.85), Inches(6.4), Inches(4.7),
            "Cascade main window \u2014 agent chat + production pipeline", dark=True)
footer(s, 3, "Overview", dark=True)

# ---------- SLIDE 4 : PART 1 DIVIDER ----------
s = slide(dark=True)
box(s, Inches(9.2), Inches(-2.2), Inches(6.5), Inches(6.5), fill=INK2, shape=MSO_SHAPE.OVAL)
box(s, Inches(11.2), Inches(4.8), Inches(3.6), Inches(3.6), fill=CORAL, shape=MSO_SHAPE.OVAL)
text(s, Inches(0.8), Inches(1.6), Inches(4), Inches(2.6), [[P("01", 150, CORAL, True)]])
text(s, Inches(0.85), Inches(4.05), Inches(10.5), Inches(1.6),
     [[P("Part One \u2014 Operations", 44, WHITE, True)],
      [P("Centralize the tools, kill the seat licenses, and put a price tag on every generation.", 17, LILAC, False)]],
     line_spacing=1.25)
chip(s, Inches(0.85), Inches(5.9), Inches(2.6), Inches(0.52), "Centralize", TEAL, INK, 13)
chip(s, Inches(3.65), Inches(5.9), Inches(2.7), Inches(0.52), "Consolidate", AMBER, INK, 13)
chip(s, Inches(6.55), Inches(5.9), Inches(2.7), Inches(0.52), "Track spend", PINK, INK, 13)
footer(s, 4, "Section divider", dark=True)

# ---------- SLIDE 5 : CENTRALIZE ----------
s = slide(dark=False)
blob_strip(s)
kicker(s, Inches(0.55), Inches(0.95), "Operations \u00b7 1 of 3", TEAL)
text(s, Inches(0.55), Inches(1.3), Inches(6.2), Inches(1.2),
     [[P("One door for every ", 30, DARKTXT, True), P("AI workflow", 30, TEAL, True)]])
bullet_list(s, Inches(0.6), Inches(2.55), Inches(5.7), [
    "One desktop app per workstation \u2014 no new web accounts, no new subscriptions.",
    "Every generation runs through Cascade with studio-owned API keys.",
    "Shared agents, skills, and style presets keep output consistent across artists.",
    "All assets save into one folder per production \u2014 nothing lost in a downloads folder.",
], dark=False, size=14, gap=0.82, bullet_color=TEAL)
box(s, Inches(0.6), Inches(6.0), Inches(5.6), Inches(0.75), fill=WHITE, line=TEAL, radius=0.5)
text(s, Inches(0.85), Inches(6.12), Inches(5.2), Inches(0.55),
     [[P("One tool to license, learn, and support.", 14, DARKTXT, True)]])
placeholder(s, Inches(6.75), Inches(1.7), Inches(6.0), Inches(4.95),
            "Pipeline / agent screen showing everything in one place", dark=False)
footer(s, 5, "Part 1 \u2014 Operations", dark=False)

# ---------- SLIDE 6 : REPLACE BOORDS ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "Operations \u00b7 2 of 3", AMBER)
text(s, Inches(0.55), Inches(0.9), Inches(11.5), Inches(1.0),
     [[P("Storyboards are built in \u2014 ", 32, WHITE, True),
       P("retire the Boords subscription", 32, AMBER, True)]], line_spacing=1.05)
box(s, Inches(0.55), Inches(2.3), Inches(3.85), Inches(4.0), fill=INK2, radius=0.08)
text(s, Inches(0.85), Inches(2.5), Inches(3.3), Inches(0.5),
     [[P("TODAY \u2014 separate subscription", 15, CORAL, True)]])
bullet_list(s, Inches(0.9), Inches(3.1), Inches(3.3), [
    "Monthly per-seat fee \u2014 for boarding only.",
    "One link in the chain; still need other tools.",
    "Style and characters rebuilt every project.",
    "Another login, vendor, and invoice.",
], dark=True, size=12.5, gap=0.72, bullet_color=CORAL)
box(s, Inches(4.6), Inches(2.3), Inches(3.85), Inches(4.0), fill=TEAL, radius=0.08)
text(s, Inches(4.9), Inches(2.5), Inches(3.3), Inches(0.5),
     [[P("WITH CASCADE \u2014 included", 15, INK, True)]])
bullet_list(s, Inches(4.95), Inches(3.1), Inches(3.3), [
    "Storyboard grid with 4-digit shot numbering.",
    "Printable PDF export \u2014 1 or 3 panels per page, with our logo.",
    "Boards via OpenArt or Higgsfield, or import your own art.",
    "Same tool builds the animatic and editor handoff.",
], dark=True, size=12.5, gap=0.72, bullet_color=INK)
placeholder(s, Inches(8.7), Inches(2.3), Inches(4.05), Inches(4.0),
            "Board grid + storyboard PDF export", dark=True)
chip(s, Inches(0.55), Inches(6.5), Inches(3.6), Inches(0.52), "One less recurring line item", AMBER, INK, 13)
footer(s, 6, "Part 1 \u2014 Operations", dark=True)

# ---------- SLIDE 7 : EXPENSE TRACKING ----------
s = slide(dark=False)
blob_strip(s)
kicker(s, Inches(0.55), Inches(0.95), "Operations \u00b7 3 of 3", PINK)
text(s, Inches(0.55), Inches(1.3), Inches(11.5), Inches(0.9),
     [[P("Every generation, ", 32, DARKTXT, True), P("on the ledger", 32, PINK, True)]])
text(s, Inches(0.55), Inches(2.15), Inches(5.6), Inches(1.4),
     [[P("Cascade keeps a running expense ledger for every AI generation made in-app \u2014 auto-priced against our own pricing rules, sortable per production, and exportable to CSV for finance.", 15, GRAY, False)]],
     line_spacing=1.15)
bullet_list(s, Inches(0.6), Inches(3.5), Inches(5.7), [
    "Automatic pricing: each image, video, or 3D generation gets a dollar value.",
    "Per-production totals \u2014 see what a project cost, not just the month.",
    "Manual rows for purchased assets (stock, music, third-party files).",
    "One-click CSV export that finance can actually read.",
], dark=False, size=13.5, gap=0.72, bullet_color=PINK)
for i, (t1, t2) in enumerate([("Auto-priced", "every generation"), ("CSV export", "finance-ready"), ("Per project", "spend by production")]):
    x = Inches(0.6) + Emu(int(Inches(1.95) * i))
    box(s, x, Inches(6.25), Inches(1.8), Inches(0.8), fill=WHITE, line=PINK, radius=0.35)
    text(s, x + Inches(0.1), Inches(6.34), Inches(1.6), Inches(0.65),
         [[P(t1, 12, DARKTXT, True)], [P(t2, 10.5, GRAY, False)]], align=PP_ALIGN.CENTER)
placeholder(s, Inches(6.75), Inches(2.05), Inches(6.0), Inches(4.6),
            "Expenses page \u2014 ledger, pricing rules, CSV export", dark=False)
footer(s, 7, "Part 1 \u2014 Operations", dark=False)

# ---------- SLIDE 8 : OPS RECAP ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "Part One recap", AMBER)
text(s, Inches(0.55), Inches(0.9), Inches(11.5), Inches(0.9),
     [[P("The business case, in three numbers", 32, WHITE, True)]])
stats = [
    ("ONE", "tool to license, learn, and support", "All AI work flows through Cascade \u2014 one install, one update cycle, one vendor relationship.", CORAL),
    ("ZERO", "new subscriptions or seat fees", "Bring Your Own API Key: we pay providers directly, with no platform markup. The Boords line item goes away.", TEAL),
    ("100%", "of AI spend, visible in one ledger", "Every generation priced and logged per production, exportable to finance in one click.", AMBER),
]
cw2 = Inches(3.95); gapx2 = Inches(0.22); y02 = Inches(2.15); ch2 = Inches(4.05)
for i, (big, mid, body, col) in enumerate(stats):
    x = Inches(0.55) + Emu(int((cw2 + gapx2) * i))
    box(s, x, y02, cw2, ch2, fill=INK2, radius=0.10)
    box(s, x, y02, cw2, Inches(0.14), fill=col, radius=0.5)
    text(s, x + Inches(0.3), y02 + Inches(0.4), cw2 - Inches(0.6), Inches(1.1),
         [[P(big, 48, col, True)]])
    text(s, x + Inches(0.3), y02 + Inches(1.5), cw2 - Inches(0.6), Inches(0.9),
         [[P(mid, 16, WHITE, True)]], line_spacing=1.1)
    text(s, x + Inches(0.3), y02 + Inches(2.45), cw2 - Inches(0.6), Inches(1.5),
         [[P(body, 12.5, LILAC, False)]], line_spacing=1.15)
footer(s, 8, "Part 1 \u2014 Operations", dark=True)

# ---------- SLIDE 9 : PART 2 DIVIDER ----------
s = slide(dark=True)
box(s, Inches(9.2), Inches(3.6), Inches(6.0), Inches(6.0), fill=INK2, shape=MSO_SHAPE.OVAL)
box(s, Inches(11.4), Inches(-1.6), Inches(3.6), Inches(3.6), fill=TEAL, shape=MSO_SHAPE.OVAL)
box(s, Inches(0.8), Inches(5.1), Inches(2.2), Inches(2.2), fill=AMBER, shape=MSO_SHAPE.OVAL)
text(s, Inches(0.8), Inches(1.6), Inches(4), Inches(2.6), [[P("02", 150, TEAL, True)]])
text(s, Inches(0.85), Inches(4.05), Inches(10.5), Inches(1.6),
     [[P("Part Two \u2014 The Creative Toolkit", 44, WHITE, True)],
      [P("One guided process from script to screen \u2014 boards, motion, animatic, and handoff in a single production file.", 17, LILAC, False)]],
     line_spacing=1.25)
chip(s, Inches(0.85), Inches(5.9), Inches(2.9), Inches(0.52), "Script " + R + " Shots", VIOLET, WHITE, 13)
chip(s, Inches(3.95), Inches(5.9), Inches(2.9), Inches(0.52), "Style " + R + " Boards", PINK, WHITE, 13)
chip(s, Inches(7.05), Inches(5.9), Inches(3.2), Inches(0.52), "Motion " + R + " Handoff", AMBER, INK, 13)
footer(s, 9, "Section divider", dark=True)

# ---------- SLIDE 10 : THE PIPELINE ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "Process unification", VIOLET)
text(s, Inches(0.55), Inches(0.9), Inches(11.5), Inches(0.9),
     [[P("Five guided steps \u2014 ", 32, WHITE, True),
       P("not a pile of chat threads", 32, VIOLET, True)]])
text(s, Inches(0.55), Inches(1.8), Inches(11.5), Inches(0.5),
     [[P("Cascade walks a production through fixed steps, saving every asset in one place. Each step feeds the next.", 14, LILAC, False)]])
steps = [
    ("1", "Script " + R + " Shots", "Ingest a script; scenes and shots broken out automatically.", CORAL),
    ("2", "Characters & Style", "Build cast looks and one master style applied to every frame.", AMBER),
    ("3", "Storyboard Boards", "Generate frames in parallel via OpenArt or Higgsfield.", TEAL),
    ("4", "Motion & Timing", "Animate key shots; assign per-shot durations.", VIOLET),
    ("5", "Handoff", "Editor handoff pack + rendered animatic via built-in ffmpeg.", PINK),
]
cw3 = Inches(2.32); gapx3 = Inches(0.13); y03 = Inches(2.6); ch3 = Inches(3.6)
for i, (n, t, b, col) in enumerate(steps):
    x = Inches(0.55) + Emu(int((cw3 + gapx3) * i))
    box(s, x, y03, cw3, ch3, fill=INK2, radius=0.10)
    chip(s, x + Inches(0.22), y03 + Inches(0.24), Inches(0.5), Inches(0.5), n, col, INK, 15)
    text(s, x + Inches(0.22), y03 + Inches(1.0), cw3 - Inches(0.44), Inches(0.9),
         [[P(t, 15, WHITE, True)]], line_spacing=1.05)
    text(s, x + Inches(0.22), y03 + Inches(1.9), cw3 - Inches(0.44), Inches(1.55),
         [[P(b, 11.5, LILAC, False)]], line_spacing=1.12)
    if i < 4:
        text(s, x + cw3 - Inches(0.03), y03 + ch3 / 2 - Inches(0.22), Inches(0.32), Inches(0.45),
             [[P(R, 16, LILAC, True)]])
footer(s, 10, "Part 2 \u2014 The Creative Toolkit", dark=True)

# ---------- SLIDE 11 : STEPS 1-2 ----------
s = slide(dark=False)
blob_strip(s)
kicker(s, Inches(0.55), Inches(0.95), "Creative toolkit \u00b7 Steps 1\u20132", CORAL)
text(s, Inches(0.55), Inches(1.3), Inches(11.5), Inches(0.9),
     [[P("From raw script to a ", 30, DARKTXT, True),
       P("numbered, styled shot list", 30, CORAL, True)]])
bullet_list(s, Inches(0.6), Inches(2.6), Inches(5.7), [
    "Drop in a script \u2014 scenes and shots broken out in one pass.",
    "4-digit shot numbers (0100 grid) leave slots for inserts, so renumbering never breaks the board.",
    "Character builder: describe the cast once; looks stay consistent across every frame.",
    "Master style sheet: one visual language stamped onto every generation.",
], dark=False, size=14, gap=0.82, bullet_color=CORAL)
box(s, Inches(0.6), Inches(6.05), Inches(5.6), Inches(0.75), fill=WHITE, line=CORAL, radius=0.5)
text(s, Inches(0.85), Inches(6.17), Inches(5.2), Inches(0.55),
     [[P("The first draft of the board lives next to the final render.", 13.5, DARKTXT, True)]])
placeholder(s, Inches(6.75), Inches(1.7), Inches(6.0), Inches(4.95),
            "Script intake + shot list / character builder", dark=False)
footer(s, 11, "Part 2 \u2014 The Creative Toolkit", dark=False)

# ---------- SLIDE 12 : STEP 3 BOARDS ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "Creative toolkit \u00b7 Step 3", TEAL)
text(s, Inches(0.55), Inches(0.9), Inches(11.5), Inches(0.9),
     [[P("Boards that ", 32, WHITE, True), P("actually match the script", 32, TEAL, True)]])
bullet_list(s, Inches(0.6), Inches(2.1), Inches(5.7), [
    "Frames generate in parallel \u2014 one click per scene, not one prompt at a time.",
    "Choose OpenArt or Higgsfield per generation; swap providers without losing the board.",
    "Character and style context is carried into every prompt automatically.",
    "Retry a single shot to fill gaps, or manually import artwork.",
], dark=True, size=14, gap=0.8, bullet_color=TEAL)
chip(s, Inches(0.6), Inches(5.55), Inches(2.75), Inches(0.52), "OpenArt \u2014 image gen", VIOLET, WHITE, 12.5)
chip(s, Inches(3.55), Inches(5.55), Inches(2.85), Inches(0.52), "Higgsfield \u2014 image + video", PINK, WHITE, 12.5)
chip(s, Inches(0.6), Inches(6.3), Inches(5.8), Inches(0.52), "Both providers, our own API keys \u2014 generation is auto-logged to the ledger", INK2, AMBER, 12.5, line=AMBER)
placeholder(s, Inches(6.75), Inches(2.0), Inches(6.0), Inches(4.65),
            "Board grid mid-generation \u2014 provider picker, retries", dark=True)
footer(s, 12, "Part 2 \u2014 The Creative Toolkit", dark=True)

# ---------- SLIDE 13 : MOTION TOOLS ----------
s = slide(dark=False)
blob_strip(s)
kicker(s, Inches(0.55), Inches(0.95), "Creative toolkit \u00b7 Step 4a", VIOLET)
text(s, Inches(0.55), Inches(1.3), Inches(11.5), Inches(0.9),
     [[P("Motion tools \u2014 ", 30, DARKTXT, True),
       P("when a still isn't enough", 30, VIOLET, True)]])
bullet_list(s, Inches(0.6), Inches(2.55), Inches(5.7), [
    "Node-graph editor: wire references, style, and prompts into image and video generation.",
    "Animate any frame \u2014 pick a video model, resolution, and clip length per shot.",
    "In-betweener: line up 2\u20135 keyframes and let it fill the motion between them.",
    "Video shots sit alongside stills on the same animatic timeline.",
], dark=False, size=14, gap=0.82, bullet_color=VIOLET)
box(s, Inches(0.6), Inches(6.05), Inches(5.6), Inches(0.75), fill=WHITE, line=VIOLET, radius=0.5)
text(s, Inches(0.85), Inches(6.17), Inches(5.2), Inches(0.55),
     [[P("Hero shots get motion without leaving the board.", 13.5, DARKTXT, True)]])
placeholder(s, Inches(6.75), Inches(1.7), Inches(6.0), Inches(4.95),
            "Node graph \u2014 references wired into image / video / in-between nodes", dark=False)
footer(s, 13, "Part 2 \u2014 The Creative Toolkit", dark=False)

# ---------- SLIDE 14 : STEPS 4-5 ANIMATIC + HANDOFF ----------
s = slide(dark=True)
kicker(s, Inches(0.55), Inches(0.55), "Creative toolkit \u00b7 Steps 4\u20135", PINK)
text(s, Inches(0.55), Inches(0.9), Inches(11.5), Inches(0.9),
     [[P("Timed animatic and a ", 32, WHITE, True),
       P("clean editor handoff", 32, PINK, True)]])
bullet_list(s, Inches(0.6), Inches(2.1), Inches(5.7), [
    "Per-shot durations assigned in one pass \u2014 total runtime locked in.",
    "Animatic rendered in-app with ffmpeg, with voiceover and music mixed in.",
    "Handoff pack built for the edit: edit decision list, After Effects script, and media manifest.",
    "Everything \u2014 script, boards, media, timing \u2014 lives in one production file.",
], dark=True, size=14, gap=0.8, bullet_color=PINK)
chip(s, Inches(0.6), Inches(5.55), Inches(2.6), Inches(0.52), "CMX3600 EDL", TEAL, INK, 12.5)
chip(s, Inches(3.4), Inches(5.55), Inches(3.0), Inches(0.52), "After Effects rebuild script", VIOLET, WHITE, 12.5)
chip(s, Inches(0.6), Inches(6.3), Inches(5.8), Inches(0.52), "Render MP4 in-app \u2014 no external tooling, no folder dumps", INK2, AMBER, 12.5, line=AMBER)
placeholder(s, Inches(6.75), Inches(2.0), Inches(6.0), Inches(4.65),
            "Animatic timeline + editor handoff / render screen", dark=True)
footer(s, 14, "Part 2 \u2014 The Creative Toolkit", dark=True)

# ---------- SLIDE 15 : CLOSING / ASK ----------
s = slide(dark=True)
box(s, Inches(10.6), Inches(-1.6), Inches(5.2), Inches(5.2), fill=VIOLET, shape=MSO_SHAPE.OVAL)
box(s, Inches(11.8), Inches(-0.7), Inches(4.0), Inches(4.0), fill=PINK, shape=MSO_SHAPE.OVAL)
box(s, Inches(-1.7), Inches(5.4), Inches(4.2), Inches(4.2), fill=TEAL, shape=MSO_SHAPE.OVAL)
text(s, Inches(0.9), Inches(1.35), Inches(8), Inches(0.4),
     [[P("THE ASK", 13, AMBER, True)]])
text(s, Inches(0.85), Inches(1.75), Inches(11), Inches(1.8),
     [[P("Adopt Cascade as ", 46, WHITE, True), P("our internal AI studio", 46, TEAL, True)]])
bullet_list(s, Inches(0.9), Inches(3.15), Inches(7.0), [
    "Pilot on one upcoming production \u2014 full pipeline, start to handoff.",
    "Provision studio API keys for OpenArt and Higgsfield (usage billed at cost).",
    "Cancel the Boords subscription at renewal.",
    "Review the expense ledger after 30 days \u2014 decide with real numbers.",
], dark=True, size=15.5, gap=0.68, bullet_color=AMBER)
chip(s, Inches(0.9), Inches(6.1), Inches(3.3), Inches(0.52), "Costs what we already spend", TEAL, INK, 13)
chip(s, Inches(4.4), Inches(6.1), Inches(3.4), Inches(0.52), "Shows us what that spend is", AMBER, INK, 13)
footer(s, 15, "Cascade \u00b7 Thank you \u2014 questions welcome", dark=True)

OUT = "Cascade_Pitch.pptx"
prs.save(OUT)
print("Saved %s with %d slides" % (OUT, len(prs.slides._sldIdLst)))
